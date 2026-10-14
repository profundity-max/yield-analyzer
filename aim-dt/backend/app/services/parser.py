from __future__ import annotations

import csv
import hashlib
import json
import math
import uuid
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from app.config import get_settings

settings = get_settings()
DEFAULT_POINTS = ["T1-T2", "T1-T3", "T1-T4", "T1-T5"]
COMMON_HEADERS = ["产品SN", "检测时间", "config"]
ERROR_SENTINELS = {-999.0, 9999.0}


@dataclass(slots=True)
class SpecCandidate:
    enabled_points: list[str]
    rules: dict[str, dict[str, dict[str, float | None]]]
    source_headers: dict[str, str]
    warnings: list[str] = field(default_factory=list)


@dataclass(slots=True)
class ParsedFile:
    events: Iterator[dict[str, Any]]
    spec: SpecCandidate | None
    row_count: int
    headers: list[str]
    warnings: list[str] = field(default_factory=list)


def point_to_column(point: str) -> str:
    return "Temp_" + point.replace("-", "_")


def normalize_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def normalize_point(value: str) -> str:
    text = normalize_header(value).upper().replace("_", "-")
    return text


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        number = float(value)
        return None if not math.isfinite(number) else number
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return None if not math.isfinite(number) else number


def parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        # Excel serial date
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).replace(tzinfo=None)
    except ValueError:
        return None


def value_token(value: Any) -> str:
    number = parse_number(value)
    if number is None:
        return "__ERROR__"
    if number in ERROR_SENTINELS:
        return "__ERROR__"
    return format(number, ".12g")


def production_day(timestamp: datetime) -> date:
    return (timestamp - timedelta(hours=settings.day_cutoff_hour)).date()


def event_hash(sn: str, timestamp: datetime, measurements: dict[str, float | None]) -> str:
    payload = {
        "sn": sn,
        "time": timestamp.isoformat(timespec="seconds"),
        "values": {point: value_token(measurements.get(point)) for point in measurements},
    }
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def event_id_from_hash(value: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, value))


def find_header_row(rows: list[tuple[Any, ...]]) -> int:
    for index, row in enumerate(rows[:30]):
        headers = {normalize_header(value) for value in row}
        if all(header in headers for header in COMMON_HEADERS):
            return index
    raise ValueError("找不到表头：至少需要 产品SN、检测时间、config 三列")


def make_header_index(header: list[Any]) -> dict[str, int]:
    return {
        normalize_header(value): index
        for index, value in enumerate(header)
        if normalize_header(value)
    }


def build_spec(
    rows: list[tuple[Any, ...]],
    header_index: dict[str, int],
    points: list[str] | None = None,
    source_headers: dict[str, str] | None = None,
) -> SpecCandidate | None:
    points = points or DEFAULT_POINTS
    source_headers = source_headers or {point: point_to_column(point) for point in points}
    unit_index = header_index.get("点位名称")
    if unit_index is None:
        unit_index = 23
    usl_row: tuple[Any, ...] | None = None
    lsl_row: tuple[Any, ...] | None = None
    for row in rows:
        marker = normalize_header(row[unit_index]) if unit_index < len(row) else ""
        if marker.upper() == "USL":
            usl_row = row
        elif marker.upper() == "LSL":
            lsl_row = row
    if usl_row is None and lsl_row is None:
        return None
    rules: dict[str, dict[str, dict[str, float | None]]] = {"default": {}, "by_config": {}}
    warnings: list[str] = []
    for point in points:
        column = source_headers.get(point, point_to_column(point))
        col_index = header_index.get(column)
        if col_index is None:
            warnings.append(f"找不到 {point} 对应的列 {column}")
            continue
        lsl = parse_number(lsl_row[col_index]) if lsl_row and col_index < len(lsl_row) else None
        usl = parse_number(usl_row[col_index]) if usl_row and col_index < len(usl_row) else None
        if lsl == -999:
            lsl = None
        if usl == 9999:
            usl = None
        rules["default"][point] = {"lsl": lsl, "usl": usl}
    return SpecCandidate(points, rules, source_headers, warnings)


def parse_xlsx(path: Path, points: list[str] | None = None) -> ParsedFile:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    buffered: list[tuple[Any, ...]] = []
    for row in worksheet.iter_rows(values_only=True):
        buffered.append(tuple(row))
        if len(buffered) >= 30:
            break
    header_row_index = find_header_row(buffered)
    header = list(buffered[header_row_index])
    header_index = make_header_index(header)
    mapped_points = points or DEFAULT_POINTS
    source_headers = {
        point: next(
            (key for key in header_index if normalize_point(key.replace("Temp_", "")) == point),
            point_to_column(point),
        )
        for point in mapped_points
    }
    spec = build_spec(buffered, header_index, mapped_points, source_headers)
    estimated_rows = max(0, (worksheet.max_row or 0) - header_row_index - 1)
    workbook.close()

    def iterator() -> Iterator[dict[str, Any]]:
        data_workbook = load_workbook(path, read_only=True, data_only=True)
        data_sheet = data_workbook.active
        try:
            for row_number, row in enumerate(data_sheet.iter_rows(values_only=True), start=1):
                if row_number <= header_row_index + 1:
                    continue
                marker_index = header_index.get("点位名称", 23)
                marker = normalize_header(row[marker_index]) if len(row) > marker_index else ""
                if marker.upper() in {"USL", "LSL", "UNIT"}:
                    continue
                sn_index = header_index["产品SN"]
                time_index = header_index["检测时间"]
                sn = normalize_header(row[sn_index]) if len(row) > sn_index else ""
                detected_at = parse_datetime(row[time_index]) if len(row) > time_index else None
                if not sn or detected_at is None:
                    continue
                measurements = {
                    point: parse_number(row[header_index[column]])
                    if header_index.get(column) is not None and len(row) > header_index[column]
                    else None
                    for point, column in source_headers.items()
                }
                digest = event_hash(sn, detected_at, measurements)
                yield {
                    "event_id": event_id_from_hash(digest),
                    "event_hash": digest,
                    "sn": sn,
                    "detected_at": detected_at,
                    "production_day": production_day(detected_at),
                    "config": normalize_header(row[header_index.get("config")])
                    if header_index.get("config") is not None and len(row) > header_index["config"]
                    else None,
                    "site_name": normalize_header(row[header_index.get("站点名称")])
                    if header_index.get("站点名称") is not None
                    and len(row) > header_index["站点名称"]
                    else None,
                    "equipment_id": normalize_header(row[header_index.get("设备编号")])
                    if header_index.get("设备编号") is not None
                    and len(row) > header_index["设备编号"]
                    else None,
                    "slot": normalize_header(row[header_index.get("Slot")])
                    if header_index.get("Slot") is not None and len(row) > header_index["Slot"]
                    else None,
                    "processing_ms": parse_number(row[header_index.get("处理时长(ms)")])
                    if header_index.get("处理时长(ms)") is not None
                    and len(row) > header_index["处理时长(ms)"]
                    else None,
                    "original_final_result": normalize_header(row[header_index.get("最终结果")])
                    if header_index.get("最终结果") is not None
                    and len(row) > header_index["最终结果"]
                    else None,
                    "measurement_json": {point: measurements[point] for point in mapped_points},
                    "source_sequence": row_number,
                    "source_row": row_number,
                }
        finally:
            data_workbook.close()

    return ParsedFile(
        iterator(),
        spec,
        estimated_rows,
        header,
        spec.warnings if spec else [],
    )


def parse_csv(path: Path, points: list[str] | None = None) -> ParsedFile:
    encoding = None
    for candidate in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=candidate, newline="") as probe:
                probe.read(8192)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if encoding is None:
        raise ValueError("CSV 编码无法识别，请使用 UTF-8 或 GB18030")

    with path.open("r", encoding=encoding, newline="") as stream:
        sample = []
        for row in csv.reader(stream):
            sample.append(tuple(row))
            if len(sample) >= 30:
                break
    header_row_index = find_header_row(sample)
    header = list(sample[header_row_index])
    header_index = make_header_index(header)
    mapped_points = points or DEFAULT_POINTS
    source_headers = {
        point: next(
            (key for key in header_index if normalize_point(key.replace("Temp_", "")) == point),
            point_to_column(point),
        )
        for point in mapped_points
    }
    spec = build_spec(sample, header_index, mapped_points, source_headers)
    with path.open("r", encoding=encoding, newline="") as stream:
        row_count = sum(1 for _ in stream) - header_row_index - 1

    def iterator() -> Iterator[dict[str, Any]]:
        with path.open("r", encoding=encoding, newline="") as stream:
            reader = csv.reader(stream)
            for row_number, row in enumerate(reader, start=1):
                if row_number <= header_row_index + 1:
                    continue
                marker_index = header_index.get("点位名称", 23)
                marker = normalize_header(row[marker_index]) if len(row) > marker_index else ""
                if marker.upper() in {"USL", "LSL", "UNIT"}:
                    continue
                sn_index = header_index["产品SN"]
                time_index = header_index["检测时间"]
                sn = normalize_header(row[sn_index]) if len(row) > sn_index else ""
                detected_at = parse_datetime(row[time_index]) if len(row) > time_index else None
                if not sn or detected_at is None:
                    continue
                measurements = {
                    point: parse_number(row[header_index[column]])
                    if header_index.get(column) is not None and len(row) > header_index[column]
                    else None
                    for point, column in source_headers.items()
                }
                digest = event_hash(sn, detected_at, measurements)
                yield {
                    "event_id": event_id_from_hash(digest),
                    "event_hash": digest,
                    "sn": sn,
                    "detected_at": detected_at,
                    "production_day": production_day(detected_at),
                    "config": normalize_header(row[header_index.get("config")])
                    if header_index.get("config") is not None and len(row) > header_index["config"]
                    else None,
                    "site_name": normalize_header(row[header_index.get("站点名称")])
                    if header_index.get("站点名称") is not None
                    and len(row) > header_index["站点名称"]
                    else None,
                    "equipment_id": normalize_header(row[header_index.get("设备编号")])
                    if header_index.get("设备编号") is not None
                    and len(row) > header_index["设备编号"]
                    else None,
                    "slot": normalize_header(row[header_index.get("Slot")])
                    if header_index.get("Slot") is not None and len(row) > header_index["Slot"]
                    else None,
                    "processing_ms": parse_number(row[header_index.get("处理时长(ms)")])
                    if header_index.get("处理时长(ms)") is not None
                    and len(row) > header_index["处理时长(ms)"]
                    else None,
                    "original_final_result": normalize_header(row[header_index.get("最终结果")])
                    if header_index.get("最终结果") is not None
                    and len(row) > header_index["最终结果"]
                    else None,
                    "measurement_json": {point: measurements[point] for point in mapped_points},
                    "source_sequence": row_number,
                    "source_row": row_number,
                }

    return ParsedFile(iterator(), spec, max(0, row_count), header, spec.warnings if spec else [])


def parse_file(path: Path, points: list[str] | None = None) -> ParsedFile:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(path, points)
    if suffix == ".csv":
        return parse_csv(path, points)
    raise ValueError("仅支持 .xlsx 和 .csv 文件")


PARQUET_SCHEMA = pa.schema(
    [
        ("event_id", pa.string()),
        ("event_hash", pa.string()),
        ("sn", pa.string()),
        ("detected_at", pa.timestamp("us")),
        ("production_day", pa.date32()),
        ("config", pa.string()),
        ("site_name", pa.string()),
        ("equipment_id", pa.string()),
        ("slot", pa.string()),
        ("processing_ms", pa.float64()),
        ("original_final_result", pa.string()),
        ("measurement_json", pa.string()),
        ("source_sequence", pa.int64()),
    ]
)


def write_parquet(
    events: Iterator[dict[str, Any]], destination: Path, batch_size: int = 50_000
) -> tuple[int, int, int]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    writer = pq.ParquetWriter(destination, PARQUET_SCHEMA, compression="zstd")
    count = 0
    error_count = 0
    duplicate_count = 0
    seen: set[str] = set()
    batch: list[dict[str, Any]] = []
    try:
        for event in events:
            digest = event["event_hash"]
            if digest in seen:
                duplicate_count += 1
                continue
            seen.add(digest)
            raw_measurement = event["measurement_json"]
            if any(
                value is None
                or (isinstance(value, (int, float)) and float(value) in ERROR_SENTINELS)
                for value in raw_measurement.values()
            ):
                error_count += 1
            event = {**event, "measurement_json": json.dumps(raw_measurement, ensure_ascii=False)}
            batch.append(event)
            if len(batch) >= batch_size:
                writer.write_table(pa.Table.from_pylist(batch, schema=PARQUET_SCHEMA))
                count += len(batch)
                batch = []
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=PARQUET_SCHEMA))
            count += len(batch)
    finally:
        writer.close()
    return count, duplicate_count, error_count
