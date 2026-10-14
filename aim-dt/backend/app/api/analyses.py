from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import duckdb
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.deps import current_user, require_role
from app.models import AnalysisRun, AuditLog, Batch, MappingVersion, SpecVersion, User
from app.schemas import AnalysisCreate, AnalysisView
from app.services.rules import canonical_json

router = APIRouter(prefix="/analyses", tags=["analyses"])


def _get_run(db: Session, run_id: str) -> AnalysisRun:
    run = db.get(AnalysisRun, run_id)
    if run is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="分析任务不存在")
    return run


def _summary(run: AnalysisRun) -> dict:
    if run.status not in {"completed", "completed_with_warnings"} or not run.summary_json:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="分析尚未完成")
    return json.loads(run.summary_json)


@router.get("", response_model=list[AnalysisView])
def list_analyses(
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[AnalysisView]:
    items = db.scalars(select(AnalysisRun).order_by(AnalysisRun.created_at.desc())).all()
    return [AnalysisView.model_validate(item) for item in items]


@router.get("/{run_id}", response_model=AnalysisView)
def get_analysis(
    run_id: str,
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> AnalysisView:
    return AnalysisView.model_validate(_get_run(db, run_id))


@router.post("", response_model=AnalysisView, status_code=status.HTTP_202_ACCEPTED)
def create_analysis(
    payload: AnalysisCreate,
    user: User = Depends(require_role("admin", "analyst")),
    db: Session = Depends(get_db),
) -> AnalysisView:
    spec = db.get(SpecVersion, payload.spec_version_id)
    if spec is None or spec.status == "retired":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="规格版本不可用")
    if payload.mapping_version_id:
        mapping = db.get(MappingVersion, payload.mapping_version_id)
        if mapping is None or mapping.status == "retired":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="Mapping 版本不可用"
            )
    batch_query = select(Batch).where(Batch.status == "ready")
    if payload.batch_ids:
        batch_query = batch_query.where(Batch.id.in_(payload.batch_ids))
    batches = list(db.scalars(batch_query))
    if not batches:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="没有可用数据批次")
    filters = {
        "batch_ids": [item.id for item in batches] if payload.batch_ids else [],
        "date_from": payload.date_from.isoformat() if payload.date_from else None,
        "date_to": payload.date_to.isoformat() if payload.date_to else None,
        "configs": payload.configs,
        "cells": payload.cells,
        "testers": payload.testers,
    }
    run = AnalysisRun(
        name=payload.name,
        filters_json=canonical_json(filters),
        spec_version_id=payload.spec_version_id,
        mapping_version_id=payload.mapping_version_id,
        requested_by=user.id,
    )
    db.add(run)
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="create_analysis",
            entity_type="analysis",
            entity_id=run.id,
            detail_json=canonical_json(filters),
        )
    )
    db.commit()
    db.refresh(run)
    return AnalysisView.model_validate(run)


@router.post("/{run_id}/cancel", response_model=AnalysisView)
def cancel_analysis(
    run_id: str,
    _user: User = Depends(require_role("admin", "analyst")),
    db: Session = Depends(get_db),
) -> AnalysisView:
    run = _get_run(db, run_id)
    if run.status not in {"queued", "running"}:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="任务当前不可取消")
    run.cancel_requested = True
    if run.status == "queued":
        run.status = "cancelled"
        run.phase = "cancelled"
    db.commit()
    db.refresh(run)
    return AnalysisView.model_validate(run)


@router.post("/{run_id}/retry", response_model=AnalysisView, status_code=status.HTTP_202_ACCEPTED)
def retry_analysis(
    run_id: str,
    _user: User = Depends(require_role("admin", "analyst")),
    db: Session = Depends(get_db),
) -> AnalysisView:
    run = _get_run(db, run_id)
    if run.status not in {"failed", "cancelled"}:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="只有失败或取消任务可以重试"
        )
    run.status = "queued"
    run.phase = "queued"
    run.progress = 0.0
    run.error_message = None
    run.cancel_requested = False
    run.started_at = None
    run.completed_at = None
    db.commit()
    db.refresh(run)
    return AnalysisView.model_validate(run)


@router.get("/{run_id}/summary")
def get_summary(
    run_id: str,
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return _summary(_get_run(db, run_id))


@router.get("/{run_id}/events")
def get_events(
    run_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    sn: str | None = None,
    outcome: str | None = None,
    cell: str | None = None,
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    run = _get_run(db, run_id)
    if run.status not in {"completed", "completed_with_warnings"} or not run.snapshot_dir:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="分析尚未完成")
    path = Path(run.snapshot_dir) / "events.parquet"
    if not path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="明细文件不存在")
    conditions: list[str] = []
    if sn:
        conditions.append("sn LIKE '" + sn.replace("'", "''") + "%'")
    if outcome:
        conditions.append("event_outcome = '" + outcome.replace("'", "''") + "'")
    if cell:
        conditions.append("cell = '" + cell.replace("'", "''") + "'")
    where = " WHERE " + " AND ".join(conditions) if conditions else ""
    path_literal = str(path).replace("'", "''")
    con = duckdb.connect(":memory:")
    try:
        total = con.execute(
            f"SELECT count(*) FROM read_parquet('{path_literal}'){where}"
        ).fetchone()[0]
        offset = (page - 1) * page_size
        rows = con.execute(
            f"""
            SELECT event_id, sn, detected_at, production_day, config, equipment_id, slot,
                   cell, tester, qpl, processing_ms, original_final_result,
                   event_outcome, error_reason, measurement_json, point_results_json
            FROM read_parquet('{path_literal}')
            {where}
            ORDER BY detected_at DESC, sn
            LIMIT {page_size} OFFSET {offset}
            """
        ).fetchall()
        names = [item[0] for item in con.description]
    finally:
        con.close()
    items = [dict(zip(names, row, strict=True)) for row in rows]
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.get("/{run_id}/exports/summary.xlsx")
def export_summary(
    run_id: str,
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    summary = _summary(_get_run(db, run_id))
    workbook = Workbook()
    ws = workbook.active
    ws.title = "KPI"
    ws.append(["指标", "值"])
    for key, value in summary["kpis"].items():
        ws.append([key, value])
    ws2 = workbook.create_sheet("按测试点")
    ws2.append(["测试点", "失败单位数", "占投入比"])
    for row in summary["pareto"]:
        ws2.append([row["point"], row["count"], row["pct_total_input"]])
    ws3 = workbook.create_sheet("按Cell")
    ws3.append(["Cell", "首次通过", "复测通过", "最终NG", "总数"])
    for row in summary["cell_yield"]:
        ws3.append(
            [row["cell"], row["first_pass"], row["retest_pass"], row["final_ng"], row["total"]]
        )
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="aim-dt-summary-{run_id}.xlsx"'}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/{run_id}/exports/details.csv")
def export_details(
    run_id: str,
    _user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    run = _get_run(db, run_id)
    if not run.snapshot_dir:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="分析尚未完成")
    path = Path(run.snapshot_dir) / "events.parquet"
    if not path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="明细文件不存在")
    path_literal = str(path).replace("'", "''")
    con = duckdb.connect(":memory:")
    cursor = con.execute(f"SELECT * FROM read_parquet('{path_literal}') ORDER BY detected_at, sn")
    columns = [item[0] for item in cursor.description]

    async def generate():
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(columns)
        yield buffer.getvalue()
        while True:
            rows = cursor.fetchmany(2000)
            if not rows:
                break
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            writer.writerows(rows)
            yield buffer.getvalue()
        con.close()

    return StreamingResponse(
        generate(),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f'attachment; filename="aim-dt-details-{run_id}.csv"'},
    )
