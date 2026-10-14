from __future__ import annotations

import argparse
import csv
import tempfile
import time
from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook

from app.services.analytics import analyze_parquet
from app.services.parser import DEFAULT_POINTS, parse_file, write_parquet


def build_csv(source: Path, destination: Path, rows: int) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    all_rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    header_row = next(
        i
        for i, row in enumerate(all_rows)
        if "产品SN" in [str(value) for value in row if value is not None]
    )
    header = all_rows[header_row]
    data_rows = [
        row for row in all_rows[header_row + 1 :] if any(value is not None for value in row)
    ]
    if not data_rows:
        raise ValueError("source workbook has no data rows")
    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        for row in all_rows[: header_row + 1]:
            writer.writerow(row)
        for index in range(rows):
            row = list(data_rows[index % len(data_rows)])
            sn_index = header.index("产品SN")
            time_index = header.index("检测时间")
            row[sn_index] = f"BENCH-{index:09d}"
            base_time = row[time_index]
            if hasattr(base_time, "year"):
                row[time_index] = (base_time + timedelta(seconds=index)).strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            writer.writerow(row)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--rows", type=int, default=100_000)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="aim-dt-bench-") as directory:
        root = Path(directory)
        csv_path = root / "input.csv"
        parquet_path = root / "events.parquet"
        t0 = time.perf_counter()
        build_csv(Path(args.source), csv_path, args.rows)
        t1 = time.perf_counter()
        parsed = parse_file(csv_path, DEFAULT_POINTS)
        count, duplicates, errors = write_parquet(parsed.events, parquet_path)
        t2 = time.perf_counter()
        spec = {
            "enabled_points": DEFAULT_POINTS,
            "rules": parsed.spec.rules
            if parsed.spec
            else {"default": {p: {"lsl": 0, "usl": 2} for p in DEFAULT_POINTS}, "by_config": {}},
        }
        summary = analyze_parquet(
            [parquet_path],
            spec=spec,
            mapping=[],
            filters={},
            snapshot_dir=root / "snapshot",
        )
        t3 = time.perf_counter()
        print(
            {
                "rows_requested": args.rows,
                "events": count,
                "duplicates": duplicates,
                "errors": errors,
                "generate_seconds": round(t1 - t0, 2),
                "parse_and_parquet_seconds": round(t2 - t1, 2),
                "analyse_seconds": round(t3 - t2, 2),
                "total_input": summary["kpis"]["total_input"],
            }
        )


if __name__ == "__main__":
    main()
