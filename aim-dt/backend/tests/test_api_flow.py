from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AnalysisRun, Batch
from app.services.jobs import process_once

SAMPLE = Path(os.environ.get("AIM_SAMPLE_XLSX", "/Users/quality/Downloads/AIM原始数据.xlsx"))


def drain_jobs() -> None:
    with SessionLocal() as db:
        while process_once(db):
            pass


@pytest.mark.skipif(not SAMPLE.exists(), reason="AIM sample workbook not available")
def test_upload_auto_analysis_mapping_and_rerun(admin_client):
    with SAMPLE.open("rb") as stream:
        response = admin_client.post(
            "/api/v1/batches",
            files={
                "file": (
                    SAMPLE.name,
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 202, response.text
    batch = response.json()
    assert batch["status"] == "uploaded"

    drain_jobs()
    batches = admin_client.get("/api/v1/batches").json()
    assert batches[0]["status"] == "ready"
    assert batches[0]["row_count"] == 182
    assert batches[0]["event_count"] == 91

    specs = admin_client.get("/api/v1/admin/specs").json()
    assert len(specs) == 2
    parsed_spec = next(item for item in specs if item["source_batch_id"] == batch["id"])
    assert parsed_spec["enabled_points_json"] == '["T1-T2","T1-T3","T1-T4","T1-T5"]'

    workbench = admin_client.get(f"/api/v1/batches/{batch['id']}/workbench").json()
    assert workbench["analysis_id"]
    assert workbench["analysis_status"] in {"completed", "completed_with_warnings"}

    summary = admin_client.get(f"/api/v1/analyses/{workbench['analysis_id']}/summary").json()
    assert summary["kpis"]["total_input"] == 90
    assert summary["kpis"]["final_ng_qty"] == 78
    assert {row["point"]: row["count"] for row in summary["pareto"]} == {
        "T1-T2": 26,
        "T1-T3": 46,
        "T1-T4": 64,
        "T1-T5": 67,
    }

    mapping = admin_client.post(
        "/api/v1/mappings/editor",
        json={
            "name": "样例 Slot1 映射",
            "batch_id": batch["id"],
            "rows": [
                {
                    "source": {
                        "site_name": "V7X_P_Delta_T",
                        "equipment_id": "L1_P_Delta_T_1@965N1",
                        "config": "V965N1",
                        "slot": "1",
                    },
                    "target": {"cell": "Cell-1", "tester": "T01", "qpl": "QPL1"},
                    "match_count": 15,
                    "status": "mapped",
                }
            ],
        },
    )
    assert mapping.status_code == 200, mapping.text
    mapping_id = mapping.json()["mapping_version_id"]
    editor = admin_client.get(
        f"/api/v1/mapping-editor?batch_id={batch['id']}&mapping_version_id={mapping_id}"
    ).json()
    assert sum(row["match_count"] for row in editor["unmapped"]) == 76

    rerun = admin_client.post(
        f"/api/v1/analyses/{workbench['analysis_id']}/rerun-with",
        json={"spec_version_id": parsed_spec["id"], "mapping_version_id": mapping_id},
    )
    assert rerun.status_code == 202, rerun.text
    drain_jobs()
    rerun_id = rerun.json()["id"]
    rerun_summary = admin_client.get(f"/api/v1/analyses/{rerun_id}/summary").json()
    assert rerun_summary["quality"]["unknown_mapping_events"] == 76
    assert rerun_summary["kpis"]["final_ng_qty"] == 78


@pytest.mark.skipif(not SAMPLE.exists(), reason="AIM sample workbook not available")
def test_rerun_accepts_retired_spec_and_mapping(admin_client):
    with SAMPLE.open("rb") as stream:
        batch = admin_client.post(
            "/api/v1/batches",
            files={"file": (SAMPLE.name, stream, "application/xlsx")},
        ).json()
    drain_jobs()
    workbench = admin_client.get(f"/api/v1/batches/{batch['id']}/workbench").json()
    first_run_id = workbench["analysis_id"]
    first_run = admin_client.get(f"/api/v1/analyses/{first_run_id}").json()
    old_spec_id = first_run["spec_version_id"]

    admin_client.post(f"/api/v1/admin/specs/{old_spec_id}/confirm")
    editor = admin_client.get(
        f"/api/v1/specs/editor?spec_id={old_spec_id}&batch_id={batch['id']}"
    ).json()
    new_spec = admin_client.post(
        "/api/v1/specs/editor",
        json={"name": "V3 重算规格", "points": editor["points"], "source_batch_id": batch["id"]},
    ).json()
    assert new_spec["id"] != old_spec_id

    old_mapping_payload = {
        "name": "历史 Mapping",
        "batch_id": batch["id"],
        "rows": [
            {
                "source": {
                    "site_name": "V7X_P_Delta_T",
                    "equipment_id": "L1_P_Delta_T_1@965N1",
                    "config": "V965N1",
                    "slot": "1",
                },
                "target": {"cell": "Cell-1", "tester": "T01", "qpl": "QPL1"},
                "match_count": 15,
                "status": "mapped",
            }
        ],
    }
    old_mapping_id = admin_client.post("/api/v1/mappings/editor", json=old_mapping_payload).json()[
        "mapping_version_id"
    ]
    new_mapping_payload = {
        **old_mapping_payload,
        "name": "最新 Mapping",
        "rows": [
            {
                **old_mapping_payload["rows"][0],
                "target": {"cell": "Cell-2", "tester": "T01", "qpl": "QPL1"},
            }
        ],
    }
    admin_client.post("/api/v1/mappings/editor", json=new_mapping_payload)

    options = admin_client.get(f"/api/v1/config/options?batch_id={batch['id']}").json()
    assert old_spec_id in {item["id"] for item in options["specs"]}
    assert old_mapping_id in {item["id"] for item in options["mappings"]}

    response = admin_client.post(
        f"/api/v1/analyses/{first_run_id}/rerun-with",
        json={"spec_version_id": old_spec_id, "mapping_version_id": old_mapping_id},
    )
    assert response.status_code == 202, response.text
    drain_jobs()
    rerun = admin_client.get(f"/api/v1/analyses/{response.json()['id']}").json()
    assert rerun["spec_version_id"] == old_spec_id
    assert rerun["mapping_version_id"] == old_mapping_id
    assert rerun["status"] in {"completed", "completed_with_warnings"}


def modified_workbook(source: Path, destination: Path, delta: float) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    sheet = workbook.active
    column = next(
        index
        for index in range(1, sheet.max_column + 1)
        if sheet.cell(1, index).value == "Temp_T1_T2"
    )
    sheet.cell(6, column).value = float(sheet.cell(6, column).value or 0) + delta
    workbook.save(destination)


@pytest.mark.skipif(not SAMPLE.exists(), reason="AIM sample workbook not available")
def test_master_module_replace_revert_and_retention(admin_client, tmp_path: Path):
    with SAMPLE.open("rb") as stream:
        admin_client.post(
            "/api/v1/batches",
            files={"file": (SAMPLE.name, stream, "application/xlsx")},
        )
    drain_jobs()
    master = admin_client.get("/api/v1/master/workbench").json()
    assert master["module_count"] == 1
    assert len(master["modules"]) == 1
    module = master["modules"][0]
    assert module["active_revision"] == 1
    assert module["previous_available"] is False

    replacement1 = tmp_path / "replacement1.xlsx"
    modified_workbook(SAMPLE, replacement1, 0.1)
    with replacement1.open("rb") as stream:
        response = admin_client.post(
            f"/api/v1/master/modules/{module['id']}/replace",
            files={"file": (replacement1.name, stream, "application/xlsx")},
        )
    assert response.status_code == 202, response.text
    drain_jobs()
    master2 = admin_client.get("/api/v1/master/workbench").json()
    module2 = master2["modules"][0]
    assert module2["active_revision"] == 2
    assert module2["previous_available"] is True
    assert module2["previous_revision"] == 1

    replacement2 = tmp_path / "replacement2.xlsx"
    modified_workbook(SAMPLE, replacement2, 0.2)
    with replacement2.open("rb") as stream:
        admin_client.post(
            f"/api/v1/master/modules/{module['id']}/replace",
            files={"file": (replacement2.name, stream, "application/xlsx")},
        )
    drain_jobs()
    master3 = admin_client.get("/api/v1/master/workbench").json()
    module3 = master3["modules"][0]
    assert module3["active_revision"] == 3
    assert module3["previous_revision"] == 2
    assert master3["module_count"] == 1

    reverted = admin_client.post(f"/api/v1/master/modules/{module['id']}/revert")
    assert reverted.status_code == 202, reverted.text
    drain_jobs()
    master4 = admin_client.get("/api/v1/master/workbench").json()
    module4 = master4["modules"][0]
    assert module4["active_revision"] == 2
    assert master4["analysis_status"] in {"completed", "completed_with_warnings"}


@pytest.mark.skipif(not SAMPLE.exists(), reason="AIM sample workbook not available")
def test_spec_and_mapping_save_refresh_master_once(admin_client, tmp_path: Path):
    with SAMPLE.open("rb") as stream:
        admin_client.post(
            "/api/v1/batches",
            files={"file": (SAMPLE.name, stream, "application/xlsx")},
        )
    drain_jobs()
    before = admin_client.get("/api/v1/master/workbench").json()
    before_runs = len(admin_client.get("/api/v1/analyses").json())

    spec = admin_client.get("/api/v1/specs/editor").json()
    spec["points"][0]["usl"] = 1.9
    admin_client.post(
        "/api/v1/specs/editor",
        json={"name": "V3 测试规格", "points": spec["points"]},
    )
    after_spec = admin_client.get("/api/v1/analyses").json()
    assert len(after_spec) == before_runs + 1
    assert after_spec[0]["status"] == "queued"

    drain_jobs()
    mapping = admin_client.get("/api/v1/mapping-editor").json()
    rows = [
        {
            **row,
            "target": row["target"]
            if row["status"] == "mapped"
            else {"cell": "Cell-1", "tester": row["source"].get("slot", "T01"), "qpl": "QPL1"},
            "status": "mapped",
        }
        for row in [*mapping["rules"], *mapping["unmapped"]]
    ]
    after_spec_runs = len(admin_client.get("/api/v1/analyses").json())
    admin_client.post(
        "/api/v1/mappings/editor",
        json={
            "name": "V3 全量映射",
            "batch_id": before["modules"][0]["active_batch_id"],
            "rows": rows,
        },
    )
    after_mapping = admin_client.get("/api/v1/analyses").json()
    assert len(after_mapping) == after_spec_runs + 1
    assert after_mapping[0]["status"] == "queued"


@pytest.mark.skipif(not SAMPLE.exists(), reason="AIM sample workbook not available")
def test_multiple_uploads_merge_into_one_master_source(admin_client, tmp_path: Path):
    first = tmp_path / "module_a.xlsx"
    second = tmp_path / "module_b.xlsx"
    modified_workbook(SAMPLE, first, 0.1)
    modified_workbook(SAMPLE, second, 0.2)
    for file in (first, second):
        with file.open("rb") as stream:
            response = admin_client.post(
                "/api/v1/batches",
                files={"file": (file.name, stream, "application/xlsx")},
            )
        assert response.status_code == 202, response.text
    drain_jobs()
    master = admin_client.get("/api/v1/master/workbench").json()
    assert master["module_count"] == 2
    assert len(master["modules"]) == 2
    assert master["total_rows"] == 364
    assert master["last_success_analysis_id"]
    with SessionLocal() as db:
        active = db.scalars(select(Batch).where(Batch.is_active.is_(True))).all()
        assert len(active) == 2
        run = db.get(AnalysisRun, master["last_success_analysis_id"])
        assert run is not None
        assert len(json.loads(run.filters_json)["batch_ids"]) == 2
